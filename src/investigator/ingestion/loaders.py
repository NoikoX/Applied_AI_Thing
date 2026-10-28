"""Turn uploaded files (PDF, DOCX, HTML, Markdown, text) and corpus records into Documents.

The parser libraries are imported inside the functions that use them: they cost about
20 MB and are only needed when someone uploads a file (memory matters on the 512 MB host).
"""

import hashlib
import io
import re
from pathlib import PurePath

from investigator.ingestion.documents import Document, Section

SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
}


class UnsupportedFormatError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def upload_doc_id(filename: str) -> str:
    """ID from the file name, so re-uploading a file with the same name updates it."""
    return "up_" + hashlib.sha256(filename.strip().lower().encode("utf-8")).hexdigest()[:16]


def load_bytes(filename: str, data: bytes) -> Document:
    ext = PurePath(filename).suffix.lower()
    fmt = SUPPORTED_EXTENSIONS.get(ext)
    if fmt is None:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFormatError(f"Unsupported file type '{ext}'. Supported: {supported}")

    parser = {"pdf": _parse_pdf, "docx": _parse_docx, "html": _parse_html}.get(fmt, _parse_text)
    title, sections, metadata = parser(data)
    sections = [s for s in sections if s.text.strip()]
    if not sections:
        raise EmptyDocumentError(f"No extractable text in '{filename}' (scanned PDFs need OCR).")

    return Document(
        doc_id=upload_doc_id(filename),
        title=title or PurePath(filename).stem,
        source=filename,
        format=fmt,
        sections=sections,
        metadata={"filename": filename, **metadata},
    )


def load_corpus_record(rec: dict) -> Document:
    """A row of data/corpus.jsonl (one HotpotQA paragraph)."""
    return Document(
        doc_id=rec["doc_id"],
        title=rec["title"],
        source=rec["source"],
        format="hotpotqa",
        sections=[Section(rec["text"])],
        sentences=rec["sentences"],
    )


def _clean(text: str) -> str:
    """Trim each line and collapse runs of blank lines into one."""
    lines = [line.strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _parse_pdf(data: bytes) -> tuple[str | None, list[Section], dict]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = [
        Section(_clean(page.extract_text() or ""), page=i)
        for i, page in enumerate(reader.pages, start=1)
    ]
    info = reader.metadata
    metadata = {"n_pages": len(reader.pages)}
    if info and info.author:
        metadata["author"] = info.author
    try:
        if info and info.creation_date:
            metadata["created"] = info.creation_date.isoformat()
    except ValueError:  # malformed date string in the PDF
        pass
    return (info.title if info else None), sections, metadata


def _parse_docx(data: bytes) -> tuple[str | None, list[Section], dict]:
    import docx

    d = docx.Document(io.BytesIO(data))
    text = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    props = d.core_properties
    metadata = {}
    if props.author:
        metadata["author"] = props.author
    if props.created:
        metadata["created"] = props.created.isoformat()
    return props.title or None, [Section(_clean(text))], metadata


def _parse_html(data: bytes) -> tuple[str | None, list[Section], dict]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript", "head"]):
        tag.decompose()
    return title, [Section(_clean(soup.get_text("\n")))], {}


def _parse_text(data: bytes) -> tuple[str | None, list[Section], dict]:
    """Plain text and Markdown. A leading '# Heading' becomes the title."""
    text = data.decode("utf-8-sig", errors="replace")
    match = re.match(r"\s*#\s+(.+)", text)
    return (match.group(1).strip() if match else None), [Section(_clean(text))], {}
